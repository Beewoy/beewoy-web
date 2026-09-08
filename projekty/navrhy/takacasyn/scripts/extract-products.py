"""Extract the legacy weekly order list into the website product catalogue."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def category_for(name: str) -> str:
    normalized = name.casefold()
    if "chlieb" in normalized or "chleb" in normalized:
        return "Chleby"
    if any(word in normalized for word in ("praclík", "tyčinky", "pagáče", "pizza")):
        return "Slané pečivo"
    if any(word in normalized for word in ("knedľa", "parené buchty")):
        return "Parené výrobky"
    if "šiška" in normalized:
        return "Smažené výrobky"
    if any(word in normalized for word in ("plundrov", "švajčiarka", "lístkové cesto", "ufo", "croisant", "croissant", "rolka", "štvorec")):
        return "Plundrované pečivo"
    if any(word in normalized for word in ("vianočka", "bábovka", "šatôčka", "buchta", "koláč", "lúpačka", "štrúdla", "makovka", "orechový rožok", "makový rožok", "pletenka sladká")):
        return "Jemné pečivo"
    if any(word in normalized for word in ("rožok", "žemla", "uzol", "bageta", "pletenka", "banketka", "granát", "venček", "sendvič", "kornspitz")):
        return "Bežné pečivo"
    return "Ostatné"


def main(source: Path, destination: Path) -> None:
    workbook = load_workbook(source, data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    products = []

    for row_number in list(range(11, 55)) + list(range(62, 112)):
        name, weight, crate = (sheet.cell(row_number, column).value for column in range(1, 4))
        if not name:
            continue
        clean_name = " ".join(str(name).split())
        clean_weight = str(weight).removesuffix(".0")
        products.append(
            {
                "id": f"p{row_number}",
                "name": clean_name,
                "weight": clean_weight,
                "crate": str(crate).removesuffix(".0"),
                "category": category_for(clean_name),
            }
        )

    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(products, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {len(products)} products to {destination}")


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]))
